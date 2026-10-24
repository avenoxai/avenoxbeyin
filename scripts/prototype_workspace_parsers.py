#!/usr/bin/env python3
"""Offline M03 packaging experiment, not a shipped parser or workspace service.

Uses candidate wheels already downloaded into the task's scratch directory. Never
installs into the running interpreter, the user's vault, or a global environment.
"""
import argparse
from email.parser import BytesParser
import hashlib
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import time
import venv
import zipfile


WORKER = r'''
import hashlib,json,sys,time
from pathlib import Path
started=time.monotonic()
from pypdf import PdfReader,PdfWriter
from docx import Document
root=Path(sys.argv[1])
if sys.argv[2]=='fixtures':
    # A tiny real text-layer PDF. PDF Unicode/font fidelity belongs to C2's suite.
    content=b'BT /F1 12 Tf 72 720 Td (Synthetic PDF course source quartz) Tj ET'
    objects=[b'<< /Type /Catalog /Pages 2 0 R >>',
      b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>',
      b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>',
      b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>',
      b'<< /Length '+str(len(content)).encode()+b' >>\nstream\n'+content+b'\nendstream']
    data=b'%PDF-1.4\n'; offsets=[0]
    for i,obj in enumerate(objects,1):
        offsets.append(len(data)); data+=str(i).encode()+b' 0 obj\n'+obj+b'\nendobj\n'
    xref=len(data);data+=b'xref\n0 6\n0000000000 65535 f \n'
    for offset in offsets[1:]:data+=('%010d 00000 n \n'%offset).encode()
    data+=b'trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n'+str(xref).encode()+b'\n%%EOF\n'
    (root/'source.pdf').write_bytes(data)
    d=Document();d.add_paragraph('Türkçe ders notu: öğrenim, ölçüm. English source quartz.')
    table=d.add_table(rows=1,cols=2);table.cell(0,0).text='Ödev';table.cell(0,1).text='Sonuç: 80 (sentetik)'
    d.save(root/'source.docx')
    blank=PdfWriter();blank.add_blank_page(width=612,height=792);blank.write(root/'scan.pdf')
    locked=PdfWriter();locked.append(root/'source.pdf');locked.encrypt('synthetic-password');locked.write(root/'locked.pdf')
    (root/'broken.pdf').write_bytes(b'invalid synthetic document')
    print(json.dumps({'status':'created'}));sys.exit(0)
result=[]
for file in sorted(root.iterdir()):
    if file.suffix not in ('.pdf','.docx'):continue
    before=hashlib.sha256(file.read_bytes()).hexdigest()
    chunks=[];status='ready'
    try:
        if file.suffix=='.pdf':
            reader=PdfReader(file)
            if reader.is_encrypted:status='requires_password'
            else:
                for page_no,page in enumerate(reader.pages,1):
                    text=page.extract_text() or ''
                    if text.strip():chunks.append({'location':{'page':page_no},'text':text})
                if not chunks:status='requires_ocr'
        else:
            document=Document(file)
            for no,block in enumerate(document.iter_inner_content(),1):
                if hasattr(block,'rows'):
                    for r,row in enumerate(block.rows,1):
                        for c,cell in enumerate(row.cells,1):
                            if cell.text:chunks.append({'location':{'block':no,'row':r,'cell':c},'text':cell.text})
                elif block.text:chunks.append({'location':{'block':no},'text':block.text})
    except Exception as exc:status='failed'
    result.append({'name':file.name,'status':status,'chunks':chunks,
                   'source_unchanged':before==hashlib.sha256(file.read_bytes()).hexdigest()})
print(json.dumps({'startup_and_extract_ms':round((time.monotonic()-started)*1000),
                  'documents':result},ensure_ascii=False))
'''


def run(arguments, *, timeout=60):
    result = subprocess.run(arguments, capture_output=True, timeout=timeout,
                            env=dict(os.environ, PYTHONDONTWRITEBYTECODE='1'))
    if result.returncode:
        raise RuntimeError(f'prototype subprocess failed (exit {result.returncode}): '
                           + result.stderr.decode('utf-8', errors='replace')[-2000:])
    return result


def prototype(wheelhouse, work):
    wheelhouse, work = Path(wheelhouse).resolve(), Path(work).resolve()
    if work.exists():
        raise ValueError('use a fresh scratch directory; existing state is preserved')
    wheels = sorted(wheelhouse.glob('*.whl'))
    if not wheels:
        raise ValueError('candidate wheelhouse required; this prototype never downloads')
    packages = []
    requirements = []
    for wheel in wheels:
        with zipfile.ZipFile(wheel) as archive:
            metadata = [name for name in archive.namelist() if name.endswith('.dist-info/METADATA')]
            if len(metadata) != 1:
                raise ValueError('wheel metadata ambiguous')
            fields = BytesParser().parsebytes(archive.read(metadata[0]))
        digest = hashlib.sha256(wheel.read_bytes()).hexdigest()
        packages.append({'name': fields['Name'], 'version': fields['Version'],
                         'wheel': wheel.name, 'sha256': digest, 'bytes': wheel.stat().st_size})
        requirements.append(f"{fields['Name']}=={fields['Version']} --hash=sha256:{digest}")
    if {item['name'].lower().replace('_', '-') for item in packages} != {
            'pypdf', 'python-docx', 'lxml', 'typing-extensions'}:
        raise ValueError('expected only the four reviewed candidate dependencies')
    work.mkdir(parents=True)
    requirements_file = work / 'candidate-requirements.txt'
    requirements_file.write_text('\n'.join(requirements) + '\n', encoding='utf-8')
    helper = work / 'Yardımcı ortam'
    started = time.monotonic()
    venv.EnvBuilder(with_pip=True).create(helper)
    python = helper / ('Scripts/python.exe' if os.name == 'nt' else 'bin/python')
    run([str(python), '-I', '-m', 'pip', '--disable-pip-version-check', 'install',
         '--no-index', '--find-links', str(wheelhouse), '--require-hashes',
         '-r', str(requirements_file)])
    setup_seconds = round(time.monotonic() - started, 3)
    fixtures = work / 'Sentetik belgeler'; fixtures.mkdir()
    worker = work / 'candidate-worker.py'; worker.write_text(WORKER, encoding='utf-8')
    run([str(python), '-I', str(worker), str(fixtures), 'fixtures'])
    extracted = json.loads(run([str(python), '-I', str(worker), str(fixtures), 'extract']).stdout)
    documents = {item['name']: item for item in extracted['documents']}
    assert 'quartz' in documents['source.pdf']['chunks'][0]['text']
    assert documents['source.pdf']['chunks'][0]['location'] == {'page': 1}
    assert any('öğrenim' in chunk['text'] for chunk in documents['source.docx']['chunks'])
    assert any(chunk['location'].get('cell') == 2 and '80' in chunk['text']
               for chunk in documents['source.docx']['chunks'])
    assert documents['scan.pdf']['status'] == 'requires_ocr'
    assert documents['locked.pdf']['status'] == 'requires_password'
    assert documents['broken.pdf']['status'] == 'failed'
    assert all(item['source_unchanged'] for item in documents.values())
    # Same-environment candidate: direct imports require third-party packages in
    # the core interpreter. Demonstrate this without touching the actual core.
    probe = "import importlib.util,json;print(json.dumps({n:importlib.util.find_spec(n) is not None for n in ['pypdf','docx','lxml']}))"
    helper_imports = json.loads(run([str(python), '-I', '-c', probe]).stdout)
    core_imports = json.loads(run([sys.executable, '-I', '-S', '-c', probe]).stdout)
    assert all(helper_imports.values()) and not any(core_imports.values())
    incomplete = work / 'incomplete-wheelhouse'; incomplete.mkdir()
    corrupt = work / 'corrupt-wheelhouse'; corrupt.mkdir()
    for wheel in wheels:
        if not wheel.name.startswith('lxml-'):
            shutil.copyfile(wheel, incomplete / wheel.name)
        shutil.copyfile(wheel, corrupt / wheel.name)
    victim = next(corrupt.glob('pypdf-*.whl'))
    victim.write_bytes(victim.read_bytes() + b'synthetic-corruption')
    faults = {}
    for name, candidate in [('missing_wheel', incomplete), ('corrupt_wheel', corrupt)]:
        fault = subprocess.run([str(python), '-I', '-m', 'pip', '--disable-pip-version-check',
            '--no-cache-dir', 'install', '--dry-run', '--ignore-installed', '--no-index',
            '--find-links', str(candidate), '--require-hashes', '-r', str(requirements_file)],
            capture_output=True, timeout=60)
        assert fault.returncode != 0, name
        faults[name] = {'exit_code': fault.returncode, 'state': 'dependency_unavailable'
                        if name == 'missing_wheel' else 'integrity_failed'}
    report = {'status': 'prototype_passed', 'platform': sys.platform,
              'python': sys.version.split()[0], 'packages': packages,
              'wheel_bytes': sum(item['bytes'] for item in packages),
              'helper_installed_bytes': sum(p.stat().st_size for p in helper.rglob('*')
                                            if p.is_file() and not p.is_symlink()),
              'offline_setup_seconds': setup_seconds, 'helper_imports': helper_imports,
              'stdlib_core_imports': core_imports, 'faults': faults, 'extraction': extracted,
              'limits': ['candidate only; not the C2 resolver or a release dependency choice',
                         'Linux CPython wheel only; no Windows/macOS or ABI matrix',
                         'text PDF English, DOCX TR/EN + table; PDF TR not verified',
                         'no OCR, crypto extras, resource limits, workspace policy or live client',
                         'offline install needs the complete verified wheelhouse; missing wheels must be surfaced']}
    (work / 'report.json').write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    return report


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--wheelhouse', type=Path, required=True)
    parser.add_argument('--work', type=Path, required=True)
    args = parser.parse_args()
    print(json.dumps(prototype(args.wheelhouse, args.work), ensure_ascii=False, indent=2))
